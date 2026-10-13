"""The investment inputs (gbpowerdata.build.investment_inputs) on synthetic workbooks laid out as the sources are: the
generation costs' sheet (technologies across, a column per commissioning year, Low/Medium/High cases, footnote markers
on names and cells, phasing rows left out, units harmonised, a changed unit refused), the Danish Energy Agency's flat
sheet (central estimate and bounds, units converted, the version from the file name), the puller's reading of that
version; and the transcribed reference table (data/reference/investment_parameters.csv): its vocabulary, units,
statuses, ranges and sources. No network needed."""
import io

import numpy as np
import openpyxl
import pandas as pd
import pytest

from gbpowerdata import config as cfg
from gbpowerdata.build import investment_inputs as I
from gbpowerdata.pull import cost_data_pull as P


def egc_workbook(path, construction_unit="£/kW"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Technical and Cost Assumptions"
    rows = [
        ["Technology", None, None, "Fixed Offshore Wind5", None, "Onshore Windr", None],
        ["Commissioning Year", None, None, 2030, "2035\n(FOAK12)", 2030, 2035],
        ["Reference plant size", None, "MW", 1000, 1000, 50, 50],
        ["Average load factor (net load factor)", None, "%", "48%11", 0.5, 0.36, "See footnote11"],
        ["Construction Cost", "Low", construction_unit, 2000, 1900, 1000, 900],
        [None, "Medium", construction_unit, 2500, 2400, 1300, 1200],
        [None, "High", construction_unit, 3000, 2900, 1500, 1400],
        ["Construction phasing", "Year 1", "%", 0.5, 0.5, 1, 1],
        ["(% spent in each year)", "Year 2", "%", 0.5, 0.5, 0, 0],
        ["Infrastructure", "Medium", "£'000", 150000, 150000, 6100, 6100],
        ["Fixed O&M", None, "£/MW/year", 46000, 43100, 18400, 16900],
        ["Hurdle Rate", "Medium", "%", 0.089, 0.089, 0.076, 0.076],
    ]
    for r in rows:
        ws.append(r)
    wb.save(path)


def test_the_generation_costs_sheet_becomes_one_row_per_technology_year_and_parameter(tmp_path):
    egc_workbook(tmp_path / "egc.xlsx")
    t = I.egc_assumptions(tmp_path / "egc.xlsx", "desnz-egc-2025", 2024)
    assert set(t["source_technology"]) == {"Fixed Offshore Wind", "Onshore Wind"}          # footnotes removed
    assert set(t["technology"]) == {"offshore_wind", "onshore_wind"}
    cap = t[(t["parameter"] == "capex_construction")].set_index(["source_technology", "year"])
    assert cap.loc[("Fixed Offshore Wind", 2030), ["value", "low", "high"]].tolist() == [2500, 2000, 3000]
    assert cap.loc[("Fixed Offshore Wind", 2035), "variant"] == "FOAK"
    assert cap["unit"].eq("GBP/kW").all() and cap["price_base_year"].eq(2024).all()
    infra = t[t["parameter"] == "capex_infrastructure"].set_index(["source_technology", "year"])["value"]
    assert infra[("Onshore Wind", 2030)] == pytest.approx(6100 / 50)                      # GBP/kW of a 50 MW plant
    fom = t[t["parameter"] == "fixed_om"].set_index(["source_technology", "year"])
    assert fom.loc[("Fixed Offshore Wind", 2030), "value"] == pytest.approx(46.0)
    assert fom.loc[("Fixed Offshore Wind", 2030), ["unit", "status"]].tolist() == ["GBP/kW/year", "converted"]
    lf = t[t["parameter"] == "load_factor"].set_index(["source_technology", "year"])
    assert lf.loc[("Fixed Offshore Wind", 2030), "value"] == pytest.approx(0.48)
    assert lf.loc[("Fixed Offshore Wind", 2030), "note"] == "footnote 11"
    assert np.isnan(lf.loc[("Onshore Wind", 2035), "value"]) and lf.loc[("Onshore Wind", 2035), "note"] == "See footnote11"
    assert "phasing" not in " ".join(t["parameter"])


def test_a_changed_unit_in_the_generation_costs_sheet_is_refused(tmp_path):
    egc_workbook(tmp_path / "egc.xlsx", construction_unit="£/MW")
    with pytest.raises(ValueError, match="expected '£/kW'"):
        I.egc_assumptions(tmp_path / "egc.xlsx", "desnz-egc-2025", 2024)


def test_the_common_technology_labels():
    assert [I.technology_of(n) for n in ("Large-scale Solar", "Solar PV <4kW", "Floating Offshore Wind",
                                          "Hydrogen CCGT (FOAK)", "Hydrogen Reciprocating Engines 5%",
                                          "CCGT + CCS Post Combustion (NOAK)", "Biomass CCS (FOAK)", "Recip Diesel 90 hrs",
                                          "OCGT 300MW 500 hr", "EfW CHP", "AD", "Mystery")] == \
        ["solar_pv", "solar_pv_small", "offshore_wind_floating", "hydrogen_turbine", "hydrogen_recip", "gas_ccus",
         "biomass_ccus", "diesel_recip", "gas_ocgt", "energy_from_waste", "anaerobic_digestion", "other"]


def dea_workbook(path):
    wb = openpyxl.Workbook()
    wb.active.title = "Index"
    wb["Index"].append(["INDEX", "Technology Data for Energy storage"])
    wb["Index"].append(["140 PTES seasonal", "Version:", "0011"])
    ws = wb.create_sheet("alldata_flat")
    ws.append(["ws", "Technology", "cat", "par", "unit", "priceyear", "note", "ref", "est", "year", "val"])
    li = ("180 Lithium Ion Battery", "Lithium-ion battery (Utility-scale)")
    ws.append([*li, "Financial data", "Specific investment [MEUR2020/MWh]", "MEUR/MWh", 2020, "J", None, "ctrl", 2030, 0.279])
    ws.append([*li, "Financial data", "Specific investment [MEUR2020/MWh]", "MEUR/MWh", 2020, "J", None, "Lower", "Uncertainty \n(2030)", 0.185])
    ws.append([*li, "Financial data", "Specific investment [MEUR2020/MWh]", "MEUR/MWh", 2020, "J", None, "Upper", "Uncertainty \n(2030)", 0.354])
    ws.append([*li, "Energy/technical data", "Round trip efficiency [%] AC", "pct.", None, "C", None, "ctrl", 2030, 91])
    ws.append([*li, "Energy/technical data", "Application", None, None, None, None, "ctrl", 2025, "System"])
    ws.append(["160 Pumped hydro storage", "Pumped hydro", "Financial data", "Fixed O&M [€/MW/year]", "EUR/MW/year", 2020, None, None, "ctrl", 2030, 9000])
    ws.append(["162 Flywheels", "Flywheel", "Financial data", "Fixed O&M [€/MW/year]", "EUR/MW/year", 2020, None, None, "ctrl", 2030, 1])
    wb.save(path)


def test_the_danish_energy_agency_sheet_converts_units_and_keeps_bounds(tmp_path):
    path = tmp_path / "dea_technology_data_energy_storage_v0011.xlsx"
    dea_workbook(path)
    t = I.dea_storage(path).set_index(["technology", "parameter"])
    spec = t.loc[("battery_lithium", "capex_specific")]
    assert [spec["value"], spec["low"], spec["high"]] == pytest.approx([279.0, 185.0, 354.0])   # EUR/kWh
    assert spec["unit"] == "EUR/kWh" and spec["price_base_year"] == 2020 and spec["status"] == "converted"
    assert "version 0011" in spec["basis"]
    assert t.loc[("battery_lithium", "round_trip_efficiency"), "value"] == pytest.approx(0.91)
    assert t.loc[("pumped_storage", "fixed_om"), ["value", "unit"]].tolist() == [pytest.approx(9.0), "EUR/kW/year"]
    assert "flywheel" not in " ".join(t.index.get_level_values(0))                             # not a sheet read


def test_the_puller_reads_the_datasheet_version(tmp_path):
    dea_workbook(tmp_path / "d.xlsx")
    assert P.dea_version((tmp_path / "d.xlsx").read_bytes()) == "0011"


def test_the_puller_takes_the_newest_atb_version_and_keeps_the_battery_rows():
    keys = ["ATB/electricity/csv/2024/", "ATB/electricity/csv/2024/v2.0.0/ATBe.csv",
            "ATB/electricity/csv/2024/v10.0.0/ATBe.csv", "ATB/electricity/csv/2024/v4.0.0/ATBe.csv"]
    assert P.atb_key(2024, keys) == ("ATB/electricity/csv/2024/v10.0.0/ATBe.csv", "v10.0.0")
    assert P.atb_key(2022, ["ATB/electricity/csv/2022/ATBe.csv"]) == ("ATB/electricity/csv/2022/ATBe.csv", "")
    header = ["technology", "core_metric_case", "crpyears", "core_metric_parameter", "value"]
    rows = [["Utility-Scale Battery Storage", "Market", "20", "CAPEX", "1"],
            ["Utility-Scale Battery Storage", "Market", "30", "CAPEX", "1"],     # another recovery period
            ["Utility-Scale Battery Storage", "R&D", "20", "CAPEX", "1"],
            ["LandbasedWind", "Market", "20", "CAPEX", "1"],
            ["Utility-Scale Battery Storage", "Exp", "20", "Fixed O&M", "1"]]
    assert P.keep(header, rows) == [rows[0], rows[4]]


def test_an_atb_edition_becomes_rows_with_its_scenarios_as_value_and_range(tmp_path):
    base = {"technology": "Utility-Scale Battery Storage", "crpyears": "20", "techdetail": "4Hr Battery Storage",
            "core_metric_variable": "2025"}
    rows = [{**base, "core_metric_case": c, "core_metric_parameter": p, "scenario": s, "units": u, "value": v}
            for c, p, s, u, v in [("Market", "CAPEX", "Moderate", "$/kW", "1600"), ("Market", "CAPEX", "Advanced", "$/kW", "1300"),
                                  ("Market", "CAPEX", "Conservative", "$/kW", "2000"), ("Exp", "CAPEX", "Moderate", "$/kW", "9"),
                                  ("Market", "Fixed O&M", "Moderate", "$/kW-yr", "40")]]
    path = tmp_path / "atb_2024_v4.0.0_utility_battery.csv"
    pd.DataFrame(rows).to_csv(path, index=False)
    t = I.atb_batteries(path).set_index("parameter")
    cap = t.loc["capex_total"]
    assert (cap["value"], cap["low"], cap["high"]) == (1600, 1300, 2000)      # "Exp" left out where "Market" exists
    assert (cap["unit"], cap["currency"], cap["price_base_year"], cap["source"]) == ("USD/kW", "USD", 2022, "nrel-atb-2024")
    assert t.loc["fixed_om", "unit"] == "USD/kW/year" and "v4.0.0" in cap["basis"]


@pytest.fixture(scope="module")
def reference():
    ref = cfg.reference_dir()
    return pd.read_csv(ref / "investment_parameters.csv"), pd.read_csv(ref / "investment_parameter_sources.csv")


def test_the_reference_table_uses_the_vocabulary_and_its_units(reference):
    t, _ = reference
    assert list(t.columns) == I.COLUMNS
    assert set(t["parameter"]) <= set(I.PARAMETERS)
    for par, unit, cur in zip(t["parameter"], t["unit"], t["currency"].fillna("")):
        want = I.PARAMETERS[par]
        assert unit == (f"{cur}{want}" if want.startswith("/") else want), (par, unit, cur)
    assert set(t["status"]) <= set(I.STATUSES)


def test_reference_values_sit_in_their_ranges_and_rates_are_plausible(reference):
    t, _ = reference
    both = t.dropna(subset=["low", "high"])
    assert ((both["low"] <= both["value"]) & (both["value"] <= both["high"])).all()
    rates = t[t["parameter"].str.startswith("hurdle_rate")]
    assert rates["value"].between(0, 0.3).all()
    money = t[t["unit"].str.match(r"^(GBP|EUR|USD)/")]
    assert money["price_base_year"].notna().all() and money["currency"].notna().all()


def test_every_cited_source_is_listed_with_its_terms(reference):
    t, src = reference
    assert set(t["source"]) <= set(src["source_id"])
    assert src[["publisher", "title", "url", "licence", "accessed"]].notna().all().all()
    assert set(I.EGC) | {"dea-storage"} <= set(src["source_id"])               # the built sources too


def test_the_loader_adds_the_reference_values_and_the_licences(tmp_path, monkeypatch):
    from gbpowerdata import processed as PR
    monkeypatch.setattr(cfg, "processed_dir", lambda: tmp_path)
    pd.DataFrame([{c: "" for c in I.COLUMNS} | {"technology": "solar_pv", "source": "desnz-egc-2025",
                                                 "parameter": "lifetime", "value": 38}]).to_csv(
        tmp_path / "investment_inputs.csv", index=False)
    t = PR.read_investment_inputs()
    assert len(t) > 100 and t.loc[t["source"] == "desnz-egc-2025", "licence"].str.contains("Open Government").all()
    assert len(PR.read_investment_inputs(reference=False)) == 1
