"""Investment inputs as dated vintages: capital and operating costs, lifetimes, efficiencies and hurdle rates of
generation and storage technologies, as each source published them -> data/processed/investment_inputs.csv.

Why: a model of entry and exit needs the costs and required returns an investor would have assumed at the time of a
decision, not today's: each edition of a source is a vintage, and its values are by commissioning (or cost) year. The
GB generation costs studies cover no batteries, so the battery costs come from the Danish Energy Agency's datasheet
here and, as transcribed values, from Mott MacDonald's 2018 study for BEIS (data/reference/investment_parameters.csv,
with CEPA's 2025 hurdle rates, which include batteries).

Sources (data/raw/cost_data/, gbpowerdata.pull.cost_data_pull; terms in data/reference/investment_parameter_sources.csv):
  egc2020_key_data_and_assumptions.xlsx, egc2023_annex_a.xlsx, egc2025_annex_a.xlsx   BEIS/DESNZ Electricity
        Generation Costs 2020, 2023 and 2025, sheet "Technical and cost assumptions" (Open Government Licence v3.0):
        every technology by commissioning year (2025-2040; 2030-2050 in 2025), in real 2018, 2021 and 2024 GBP.
  dea_technology_data_energy_storage_v*.xlsx   Danish Energy Agency, Technology Data for Energy Storage, sheet
        "alldata_flat" (CC BY 4.0): pumped hydro (sheet 160), lithium-ion (180) and vanadium redox flow (181)
        batteries, central estimates by year with lower and upper bounds for 2030 and 2050, in 2020 EUR.
  atb_{edition}[_v*]_utility_battery.csv   NREL's Annual Technology Baseline, editions 2021-2025 (CC BY 4.0): US
        utility-scale lithium-ion batteries of 2-10 hours, CAPEX (with OCC, the overnight cost, and GCC, the grid
        connection from 2024), fixed and variable O&M (round-trip efficiency in 2025), by year, the Moderate scenario
        with the Advanced and Conservative ones as the range, in the edition's dollar year (ATB_DOLLAR_YEAR).

Rules:
  - Each source keeps its own price base (price_base_year, currency): no deflator or exchange rate is applied, so
    vintages compare only after the user converts them.
  - The ATB is US costs, kept for its yearly vintages of battery costs (2019-2023 dollars); the market case
    ("Market", "Exp" in 2025, where "Market" is missing) with a 20-year cost recovery period.
  - Units are harmonised to per kW and per kWh: GBP/MW/year to GBP/kW/year (x 0.001), the generation costs'
    infrastructure (GBP thousand for the reference plant) to GBP/kW by the plant's size, MEUR/MW to EUR/kW and
    MEUR/MWh to EUR/kWh (x 1000), percentages to fractions; status "converted" says a value was, "read" that it was
    taken as published.
  - The generation costs' Low, Medium and High cases are low, value and high of one row; the Danish Energy Agency's
    central estimate is value, its Lower and Upper bounds (2030 and 2050 only) low and high.
  - Technologies keep the source's name (source_technology, footnote markers removed) and get a common technology
    label (TECHNOLOGIES rules: solar_pv, onshore_wind, offshore_wind, gas_ccgt, battery_lithium, pumped_storage, ...).
  - Phasing rows (the share of spending in each year) are left out.
Read the table with gbpowerdata.processed.read_investment_inputs(), which adds the transcribed reference values
(CEPA's and Mott MacDonald's, cited with locators and not packed) unless reference=False.

Output: data/processed/investment_inputs.csv: technology, source_technology, variant, year, parameter, value, low,
high, unit, currency, price_base_year, basis, source, locator, status, note.

Usage: python -m gbpowerdata.build.investment_inputs
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

import numpy as np
import pandas as pd

from gbpowerdata import config as cfg

COLUMNS = ["technology", "source_technology", "variant", "year", "parameter", "value", "low", "high", "unit",
           "currency", "price_base_year", "basis", "source", "locator", "status", "note"]
STATUSES = ("read", "converted", "transcribed", "derived")
# parameter -> unit after harmonisation
PARAMETERS = {
    "plant_size": "MW", "energy_capacity": "MWh", "load_factor": "fraction", "efficiency": "fraction",
    "round_trip_efficiency": "fraction", "round_trip_efficiency_dc": "fraction", "availability": "fraction",
    "forced_outage": "fraction", "self_discharge": "fraction/day", "usable_depth_of_discharge": "fraction",
    "lifetime": "years", "cycle_life": "cycles", "predevelopment_period": "years", "construction_period": "years",
    "capex_predevelopment": "/kW", "capex_construction": "/kW", "capex_infrastructure": "/kW", "capex_total": "/kW",
    "capex_overnight": "/kW", "capex_grid_connection": "/kW",
    "capex_pump": "/kW", "capex_power": "/kW", "capex_energy": "/kWh", "capex_specific": "/kWh", "capex_other": "/kWh",
    "fixed_om": "/kW/year", "variable_om": "/MWh", "insurance": "/kW/year", "connection_uos": "/kW/year",
    "decommissioning": "/kW", "co2_transport_storage": "/tCO2", "hurdle_rate": "fraction",
    "hurdle_rate_unlevered": "fraction", "hurdle_rate_merchant_uplift": "fraction",
}
EGC = {  # source id -> (file, price base year)
    "desnz-egc-2020": ("egc2020_key_data_and_assumptions.xlsx", 2018),
    "desnz-egc-2023": ("egc2023_annex_a.xlsx", 2021),
    "desnz-egc-2025": ("egc2025_annex_a.xlsx", 2024),
}
# generation costs row label (start, lower case) -> (parameter, the source unit it must carry, factor to the unit)
EGC_ROWS = {
    "reference plant size": ("plant_size", "MW", 1.0),
    "average load factor": ("load_factor", "%", 1.0),
    "average fuel efficiency": ("efficiency", "%", 1.0),
    "operating lifetime": ("lifetime", "years", 1.0),
    "pre-development cost": ("capex_predevelopment", "£/kW", 1.0),
    "pre-development period": ("predevelopment_period", "years", 1.0),
    "construction cost": ("capex_construction", "£/kW", 1.0),
    "construction period": ("construction_period", "years", 1.0),
    "infrastructure": ("capex_infrastructure", "£'000", None),           # per reference plant: divided by its size
    "fixed o&m": ("fixed_om", "£/MW/year", 0.001),
    "variable o&m": ("variable_om", "£/MWh", 1.0),
    "insurance": ("insurance", "£/MW/year", 0.001),
    "connection and use of system charges": ("connection_uos", "£/MW/year", 0.001),
    "hurdle rate": ("hurdle_rate", "%", 1.0),
    "co2 capture and storage": ("co2_transport_storage", "£/tCO2", 1.0),
    "decommissioning cost": ("decommissioning", "£/MW", 0.001),
}
EGC_SKIP = ("pre-development phasing", "construction phasing", "technology", "commissioning year")
# footnote letters the 2023 edition leaves on three names, after its footnote numbers are removed
EGC_NAMES = {"Onshore Windr": "Onshore Wind", "Floating Offshore Windr": "Floating Offshore Wind",
             "Tidal Stream Energyr": "Tidal Stream Energy"}
DEA_SHEETS = {"160 Pumped hydro storage": "pumped_storage", "180 Lithium Ion Battery": "battery_lithium",
              "181 Vanadium Redox Flow Battery": "battery_flow"}
DEA_ROWS = [  # (regex on the parameter text, parameter)
    (r"^Round trip efficiency \[%\] AC", "round_trip_efficiency"),
    (r"^Total efficiency \[%\] net", "round_trip_efficiency"),
    (r"^Round trip efficiency( -)? ?(\[%\] )?DC", "round_trip_efficiency_dc"),
    (r"^Technical lifetime", "lifetime"), (r"^Construction time", "construction_period"),
    (r"^Lifetime in total number of cycles", "cycle_life"),
    (r"^(Output power|Output capacity|Generating capacity) for one unit", "plant_size"),
    (r"^Energy storage capacity for one unit", "energy_capacity"), (r"^Forced outage", "forced_outage"),
    (r"^Energy losses during storage", "self_discharge"),
    (r"^Specific investment", "capex_specific"), (r"^- energy component", "capex_energy"),
    (r"^- (power|capacity) component", "capex_power"), (r"^- other project costs", "capex_other"),
    (r"^Investment, total, greenfield", "capex_total"), (r"^Investment, pump part", "capex_pump"),
    (r"^Fixed O&M", "fixed_om"), (r"^Variable O&M", "variable_om"),
]
DEA_UNITS = {  # the datasheet's unit -> (factor, unit)
    "pct.": (0.01, "fraction"), "pct./day": (0.01, "fraction/day"), "MEUR/MWh": (1000.0, "EUR/kWh"),
    "MEUR/MW": (1000.0, "EUR/kW"), "kEUR/MW/year": (1.0, "EUR/kW/year"), "EUR/MW/year": (0.001, "EUR/kW/year"),
    "EUR/MWh": (1.0, "EUR/MWh"), "kEUR/MWh": (1000.0, "EUR/MWh"), "years": (1.0, "years"), "MW": (1.0, "MW"),
    "MWh": (1.0, "MWh"), "": (1.0, "cycles"),
}
ATB_DOLLAR_YEAR = {2021: 2019, 2022: 2020, 2023: 2021, 2024: 2022, 2025: 2023}   # each edition's Definitions page
ATB_ROWS = {"CAPEX": "capex_total", "OCC": "capex_overnight", "GCC": "capex_grid_connection", "Fixed O&M": "fixed_om",
            "Variable O&M": "variable_om", "Round-Trip Efficiency": "round_trip_efficiency"}
ATB_UNITS = {"$/kW": "USD/kW", "$/kW-yr": "USD/kW/year", "$/KW-yr": "USD/kW/year", "$/MWh": "USD/MWh", "": "fraction"}
ATB_CASES = {"Moderate": "value", "Advanced": "low", "Conservative": "high"}
TECHNOLOGIES = [  # (regex on the lower-case source name, technology), first match wins
    (r"floating offshore wind", "offshore_wind_floating"), (r"offshore wind", "offshore_wind"),
    (r"onshore wind", "onshore_wind"), (r"large[- ]scale solar", "solar_pv"), (r"solar pv", "solar_pv_small"),
    (r"ccus|ccs", "gas_ccus"), (r"hydrogen recip", "hydrogen_recip"), (r"hydrogen|ccht|ocht", "hydrogen_turbine"),
    (r"ccgt", "gas_ccgt"),
    (r"ocgt", "gas_ocgt"), (r"recip.*diesel", "diesel_recip"), (r"recip", "gas_recip"),
    (r"biomass chp", "biomass_chp"), (r"biomass", "biomass"), (r"^efw", "energy_from_waste"),
    (r"^ad\b", "anaerobic_digestion"), (r"^act\b", "act"), (r"landfill", "landfill_gas"), (r"sewage", "sewage_gas"),
    (r"geothermal", "geothermal"), (r"hydro", "hydropower"), (r"wave", "wave"), (r"tidal stream", "tidal_stream"),
]


def technology_of(name: str) -> str:
    n = str(name).lower()
    if "ccs" in n or "ccus" in n:
        return "biomass_ccus" if "biomass" in n else "gas_ccus"
    return next((t for rx, t in TECHNOLOGIES if re.search(rx, n)), "other")


# ---------------------------------------------------------------------------- generation costs
def clean_name(cell) -> str:
    """A technology name without its footnote markers: "Fixed Offshore Wind5" -> "Fixed Offshore Wind",
    "Floating Offshore Wind (FOAK)5,12" -> "Floating Offshore Wind (FOAK)", "Onshore Windr" -> "Onshore Wind"."""
    s = re.sub(r"\s+", " ", str(cell)).strip()
    s = re.sub(r"(?<=[A-Za-z)])\d+(,\d+)*$", "", s)
    return EGC_NAMES.get(s, s)


def year_cell(cell) -> tuple[int | None, str]:
    """(commissioning year, note) of a year header: "2030\\n(FOAK12)" -> (2030, "FOAK")."""
    m = re.match(r"\s*(\d{4})\s*(?:\((.*)\))?", str(cell), re.S)
    if not m:
        return None, ""
    note = re.sub(r"\b(FOAK|NOAK|Demo)\d+", r"\1", m.group(2) or "").strip()
    return int(m.group(1)), note


def number(cell) -> tuple[float, str]:
    """A cell's number (a percentage string as a fraction) and the footnote it carried: "48%11" -> (0.48, "11")."""
    if cell is None or cell == "":
        return np.nan, ""
    if isinstance(cell, (int, float)):
        return float(cell), ""
    m = re.match(r"\s*(-?[\d.]+)\s*(%?)\s*(\d*)\s*$", str(cell))
    if not m:
        return np.nan, str(cell).strip()
    v = float(m.group(1)) / (100.0 if m.group(2) else 1.0)
    return v, m.group(3)


def egc_assumptions(path: Path, source: str, price_base: int) -> pd.DataFrame:
    """The "Technical and cost assumptions" sheet of a generation costs edition as rows of COLUMNS."""
    import openpyxl
    from openpyxl.utils import get_column_letter

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = next(n for n in wb.sheetnames if n.lower().startswith("technical and cost"))
    grid = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
    width = max(len(r) for r in grid)
    grid = [r + [None] * (width - len(r)) for r in grid]
    techs, cur = [], None
    for c in grid[0]:
        cur = clean_name(c) if c not in (None, "") else cur
        techs.append(cur)
    years = [year_cell(c) for c in grid[1]]
    cols = [j for j in range(3, width) if techs[j] and years[j][0]]
    size_row = next(i for i, r in enumerate(grid) if str(r[0]).lower().startswith("reference plant size"))
    found: dict[tuple, dict] = {}
    label = None
    for i, r in enumerate(grid[2:], start=2):
        if r[0] not in (None, "") and not str(r[0]).startswith("("):
            label = re.sub(r"\s+", " ", str(r[0])).strip().lower()
        if label is None or label.startswith(EGC_SKIP):
            continue
        key = next((k for k in EGC_ROWS if label.startswith(k)), None)
        if key is None:
            continue
        par, unit, factor = EGC_ROWS[key]
        if str(r[2]).strip() != unit:
            raise ValueError(f"{path.name}: row {i + 1} '{label}' in '{r[2]}', expected '{unit}'")
        case = str(r[1]).strip().lower() if r[1] not in (None, "") else "value"
        for j in cols:
            v, foot = number(r[j])
            if np.isnan(v) and not foot:
                continue
            if factor is None:                              # GBP thousand for the reference plant -> GBP/kW
                mw, _ = number(grid[size_row][j])
                v, status = v * 1000.0 / (mw * 1000.0) if mw else np.nan, "converted"
            else:                                           # the sheets hold percentages as fractions already
                v, status = v * factor, ("converted" if factor != 1.0 else "read")
            rec = found.setdefault((techs[j], years[j][0], years[j][1], par), {
                "technology": technology_of(techs[j]), "source_technology": techs[j], "variant": years[j][1],
                "year": years[j][0], "parameter": par, "value": np.nan, "low": np.nan, "high": np.nan,
                "unit": unit_of(par, "GBP"), "currency": "GBP" if PARAMETERS[par].startswith("/") else "",
                "price_base_year": price_base if PARAMETERS[par].startswith("/") else np.nan,
                "basis": f"commissioning year; real {price_base} GBP; low, value and high are the Low, Medium and High cases",
                "source": source, "locator": f"sheet '{sheet}', column {get_column_letter(j + 1)}", "status": status,
                "note": ""})
            slot = {"low": "low", "medium": "value", "high": "high", "value": "value"}.get(case, "value")
            rec[slot] = v
            rec["locator"] += f", row {i + 1}" if f"row {i + 1}" not in rec["locator"] else ""
            if foot:
                rec["note"] = f"footnote {foot}" if foot.isdigit() else foot
    return pd.DataFrame(list(found.values()), columns=COLUMNS)


def unit_of(par: str, currency: str) -> str:
    u = PARAMETERS[par]
    return f"{currency}{u}" if u.startswith("/") else u


# ---------------------------------------------------------------------------- Danish Energy Agency
def dea_storage(path: Path, source: str = "dea-storage") -> pd.DataFrame:
    """Pumped hydro, lithium-ion and vanadium redox flow batteries from the datasheet's flat sheet, as rows of
    COLUMNS."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb["alldata_flat"].iter_rows(values_only=True))
    d = pd.DataFrame(rows[1:], columns=rows[0])
    d = d[d["ws"].isin(DEA_SHEETS)]
    version = re.search(r"_v([^_]+)\.xlsx$", path.name)
    out: dict[tuple, dict] = {}
    for r in d.itertuples(index=False):
        par = next((p for rx, p in DEA_ROWS if re.search(rx, str(r.par))), None)
        if par is None:
            continue
        unit = "" if r.unit is None or (isinstance(r.unit, float) and np.isnan(r.unit)) else str(r.unit)
        factor, to = DEA_UNITS[unit]
        try:
            v = float(r.val) * factor
        except (TypeError, ValueError):
            continue
        m = re.search(r"(\d{4})", str(r.year))
        if not m:
            continue
        year = int(m.group(1))
        slot = {"ctrl": "value", "Lower": "low", "Upper": "high"}[str(r.est)]
        money = to.startswith(("EUR",))
        rec = out.setdefault((r.ws, year, par), {
            "technology": DEA_SHEETS[r.ws], "source_technology": r.Technology, "variant": "", "year": year,
            "parameter": par, "value": np.nan, "low": np.nan, "high": np.nan, "unit": to,
            "currency": "EUR" if money else "", "price_base_year": float(r.priceyear) if money and pd.notna(r.priceyear) else np.nan,
            "basis": "Danish Energy Agency technology data" + (f" version {version.group(1)}" if version else "")
                     + ": central estimate; low and high its lower and upper bounds (2030 and 2050)",
            "source": source, "locator": f"sheet 'alldata_flat': '{r.ws}', '{r.par}'",
            "status": "converted" if factor != 1.0 else "read", "note": ""})
        rec[slot] = v
    return pd.DataFrame(list(out.values()), columns=COLUMNS)


# ---------------------------------------------------------------------------- NREL ATB
def atb_batteries(path: Path) -> pd.DataFrame:
    """A cut of an ATB edition (gbpowerdata.pull.cost_data_pull: utility-scale batteries, market case, 20-year cost
    recovery) as rows of COLUMNS: one per duration (variant), year and parameter; value the Moderate scenario, low
    and high the Advanced and Conservative ones."""
    m = re.match(r"atb_(\d{4})(?:_(v[\d.]+))?_utility_battery\.csv$", path.name)
    edition, version = int(m.group(1)), m.group(2) or ""
    t = pd.read_csv(path, dtype=str, keep_default_na=False)
    t = t[t["core_metric_parameter"].isin(ATB_ROWS) & t["scenario"].isin(ATB_CASES)]
    if (t["core_metric_case"] == "Market").any():          # "Exp" stands in only where an edition has no "Market"
        t = t[t["core_metric_case"] == "Market"]
    out: dict[tuple, dict] = {}
    for r in t.itertuples(index=False):
        par = ATB_ROWS[r.core_metric_parameter]
        unit = ATB_UNITS[r.units]
        money = unit.startswith("USD")
        rec = out.setdefault((r.techdetail, int(r.core_metric_variable), par), {
            "technology": "battery_lithium", "source_technology": r.technology, "variant": r.techdetail,
            "year": int(r.core_metric_variable), "parameter": par, "value": np.nan, "low": np.nan, "high": np.nan,
            "unit": unit, "currency": "USD" if money else "", "price_base_year": ATB_DOLLAR_YEAR[edition] if money else np.nan,
            "basis": f"NREL ATB {edition} {version}".strip() + f": case '{r.core_metric_case}', 20-year cost recovery; "
                     "value, low and high the Moderate, Advanced and Conservative scenarios",
            "source": f"nrel-atb-{edition}", "locator": f"ATBe.csv, technology 'Utility-Scale Battery Storage', "
                                                         f"core_metric_parameter '{r.core_metric_parameter}'",
            "status": "read", "note": ""})
        rec[ATB_CASES[r.scenario]] = float(r.value)
    return pd.DataFrame(list(out.values()), columns=COLUMNS)


# ---------------------------------------------------------------------------- build
def build(raw: Path | None = None) -> pd.DataFrame:
    raw = (cfg.raw_dir() if raw is None else Path(raw)) / "cost_data"
    parts = []
    for source, (name, base) in EGC.items():
        if (raw / name).exists():
            parts.append(egc_assumptions(raw / name, source, base))
        else:
            print(f"  {name} missing: python -m gbpowerdata.pull.cost_data_pull", flush=True)
    dea = sorted(raw.glob("dea_technology_data_energy_storage_v*.xlsx"))
    if dea:
        parts.append(dea_storage(dea[-1]))
    else:
        print("  no dea_technology_data_energy_storage_v*.xlsx: python -m gbpowerdata.pull.cost_data_pull dea_storage")
    atb = {}
    for p in sorted(raw.glob("atb_*_utility_battery.csv")):     # the newest version of each edition
        atb[int(re.match(r"atb_(\d{4})", p.name).group(1))] = p
    parts += [atb_batteries(p) for _, p in sorted(atb.items())]
    if not parts:
        raise SystemExit("investment_inputs: no input in data/raw/cost_data/")
    t = pd.concat(parts, ignore_index=True)[COLUMNS]
    return t.sort_values(["technology", "source", "source_technology", "year", "parameter"]).reset_index(drop=True)


def main(argv: list[str] | None = None) -> int:
    argparse.ArgumentParser(description=__doc__.split("\n")[0]).parse_args(argv)
    t = build()
    path = cfg.processed_dir() / "investment_inputs.csv"
    t.to_csv(path, index=False)
    print(f"  {path.name}: {len(t):,} rows; {t.groupby('source')['source_technology'].nunique().to_dict()} technologies")
    return 0


if __name__ == "__main__":
    sys.exit(main())
