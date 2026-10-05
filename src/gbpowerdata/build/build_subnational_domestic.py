"""Extract the GB domestic-electricity series from the DESNZ subnational workbook.

Source (manual download, data/raw/manual/desnz_subnational/):
  "Subnational electricity consumption statistics 2005 to 2024", DESNZ National
  Statistics, released 18 Dec 2025.
  https://www.gov.uk/government/collections/sub-national-electricity-consumption-data
  file: Subnational_electricity_consumption_statistics_2005-2024.xlsx

Scope: GREAT BRITAIN (England, Scotland, Wales); Northern Ireland EXCLUDED — the
correct GB scope for this model (DUKES/ECUK are UK-wide by contrast).

What it closes (bill-delta feasibility gaps #1 and #2):
  - domestic SHARE of GB metered electricity consumption, per year (was a flat ~39%
    proxy from supplier segmental statements; the real series varies 37.3-42.2%).
  - GB DOMESTIC electricity METER count, per year (was a hardcoded 28.5m).

The GB total is the row coded K03000001 ("Great Britain (inc unallocated)",
"All local authorities") on each year's worksheet. Columns are keyed by their
(newline-containing) header text, matched normalised so the extractor is robust to
column re-ordering between years.

NOTE on the denominator: "Total consumption (GWh): All meters" is metered final
consumption (~249 TWh GB in 2024), NOT gross supply/demand (~316-323 TWh UK in
DUKES, which adds losses/own-use/exports). The domestic SHARE here is the correct
like-for-like match to the segmental-statement proxy, but if the model smears a
policy cost over a *levy-eligible supply* base (e.g. consumer_levy_annual.csv
eligible_demand_twh ~264 TWh) the share and the base must be reconciled before use.

Run: python -m gbpowerdata.build.build_subnational_domestic
"""

from __future__ import annotations

import re

import openpyxl
import pandas as pd

from gbpowerdata import config as cfg

SRC = (cfg.raw_dir() / "manual" / "desnz_subnational"
       / "Subnational_electricity_consumption_statistics_2005-2024.xlsx")
OUT = cfg.processed_dir() / "domestic_supply_by_year.csv"

GB_CODE = "K03000001"


def _norm(s) -> str:
    """Lower-case, strip, collapse all whitespace (incl. header newlines)."""
    return re.sub(r"\s+", " ", str(s).strip().lower()) if s is not None else ""


# normalised header text -> output column name
WANT = {
    "number of meters (thousands): all domestic": "domestic_meters_thousands",
    "number of meters (thousands): all meters": "all_meters_thousands",
    "total consumption (gwh): all domestic": "domestic_gwh",
    "total consumption (gwh): all non-domestic": "nondomestic_gwh",
    "total consumption (gwh): all meters": "total_gwh",
}


def _extract_year(ws) -> dict | None:
    """Pull the GB (K03000001) row from one year sheet, keyed by header text."""
    header_row = None
    header_map = {}
    for r, row in enumerate(ws.iter_rows(min_row=1, max_row=12, values_only=True), start=1):
        if row and _norm(row[0]) == "code":
            header_row = r
            for ci, cell in enumerate(row):
                key = _norm(cell)
                if key in WANT:
                    header_map[WANT[key]] = ci
            break
    if header_row is None or len(header_map) < len(WANT):
        return None
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if row and row[0] == GB_CODE:
            return {name: row[ci] for name, ci in header_map.items()}
    return None


def main() -> None:
    if not SRC.exists():
        raise SystemExit(
            f"Missing source workbook: {SRC}\n"
            "Download it (see the module docstring) into data/raw/manual/desnz_subnational/."
        )
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    rows = []
    for name in wb.sheetnames:
        if not re.fullmatch(r"\d{4}", name):
            continue
        rec = _extract_year(wb[name])
        if rec is None:
            continue
        rec["year"] = int(name)
        rec["domestic_share"] = rec["domestic_gwh"] / rec["total_gwh"]
        rows.append(rec)
    wb.close()

    df = (pd.DataFrame(rows)
          .sort_values("year")
          [["year", "domestic_gwh", "nondomestic_gwh", "total_gwh", "domestic_share",
            "domestic_meters_thousands", "all_meters_thousands"]]
          .reset_index(drop=True))
    # round to sensible precision (GWh to 1 dp, share to 4 dp, meters to 3 dp)
    df["domestic_gwh"] = df["domestic_gwh"].round(1)
    df["nondomestic_gwh"] = df["nondomestic_gwh"].round(1)
    df["total_gwh"] = df["total_gwh"].round(1)
    df["domestic_share"] = df["domestic_share"].round(4)
    df["domestic_meters_thousands"] = df["domestic_meters_thousands"].round(3)
    df["all_meters_thousands"] = df["all_meters_thousands"].round(3)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(OUT, index=False)
    print(f"Wrote {OUT}  ({len(df)} years: {df.year.min()}-{df.year.max()})")
    show = df[df.year >= 2016].copy()
    show["domestic_share_pct"] = (100 * show["domestic_share"]).round(1)
    show["domestic_meters_m"] = (show["domestic_meters_thousands"] / 1000).round(3)
    print(show[["year", "domestic_gwh", "total_gwh", "domestic_share_pct",
                "domestic_meters_m"]].to_string(index=False))


if __name__ == "__main__":
    main()
